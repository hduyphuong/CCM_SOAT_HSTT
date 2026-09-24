"""Phân loại + tự kiểm 4 lớp 1 HSTT so với FILE KHUNG CCM v3 (chỉ ĐỌC khung).
Cờ: CHAN (khoá nút Đồng ý) · LUU_Y (anh xem, vẫn ghi sổ được) · mỗi cờ có lớp, vị trí, số HSTT, số đối chiếu."""
import openpyxl, re, warnings
from collections import defaultdict
from doc_hstt import na, num
warnings.filterwarnings("ignore")
NGUONG = 10          # đồng — lệch ≤ 10đ coi là khớp (anh chốt)

# ---- vị trí cột khung v3 (khớp tao_khung_v3.py) — kiểm tiêu đề khi mở để bắt khung lạ
N6 = dict(ma_hd="A", loai_ghi="B", ma_doi_tac="C", ma_goi="D", so_hd="E", gia_tri="I", vat="J", pct_tu="K", pct_tt_dot="L")
N7 = dict(ma_hd="A", stt="B", pham_vi="D", noi_dung="E", dvt="F", kl_hd="G", don_gia="H", nhom="K")
N9 = dict(ma_hd="A", dot="B", ngay="C", loai="D", stt="E", kl="H", dg="I", so_tien="J")
TIEU_DE = {("N6_HD_DoiTac", "E"): "Số hợp đồng", ("N7_HD_DoiTac_ChiTiet", "E"): "Nội dung", ("N9_TT_DoiTac", "H"): "KL kỳ này"}

def _col(ws, letter, r0=2):
    return [c for c in ws[f"{letter}{r0}:{letter}{ws.max_row}"]]

def doc_khung(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    for (sh, col), lab in TIEU_DE.items():
        v = wb[sh][f"{col}1"].value
        if v != lab: raise ValueError(f"Khung không đúng mẫu v3: {sh}!{col}1 = {v!r}, cần {lab!r}")
    k = dict(path=path, hd={}, dong={}, lk_kl=defaultdict(float), lk_tien=defaultdict(float), dot_cuoi={}, tu_treo=defaultdict(float),
             dot_cuoi_tu={}, doi_tac={}, nhom_ns={}, ns=defaultdict(float), th_ns=defaultdict(float), sheet_rows={}, lk_tien_dong=defaultdict(float))
    w1 = wb["N1_DanhMuc"]
    for r in range(3, w1.max_row + 1):
        if w1[f"G{r}"].value: k["doi_tac"][w1[f"G{r}"].value] = (w1[f"H{r}"].value or "", w1[f"I{r}"].value or "")
        if w1[f"Q{r}"].value: k["nhom_ns"][w1[f"Q{r}"].value] = w1[f"T{r}"].value or ""
    w = wb["N6_HD_DoiTac"]
    for r in range(2, w.max_row + 1):
        m = w[f"A{r}"].value
        if not m: continue
        rec = k["hd"].setdefault(m, dict(ma_hd=m, ma_doi_tac=w[f"C{r}"].value, ma_goi=w[f"D{r}"].value, so_hd=str(w[f"E{r}"].value or ""),
                                         gia_tri=0.0, vat=num(w[f"J{r}"].value), pct_tu=num(w[f"K{r}"].value), pct_tt_dot=w[f"L{r}"].value))
        rec["gia_tri"] += num(w[f"I{r}"].value)
    w = wb["N7_HD_DoiTac_ChiTiet"]
    for r in range(2, w.max_row + 1):
        m = w[f"A{r}"].value
        if not m: continue
        k["dong"].setdefault(m, []).append(dict(row=r, stt=str(w[f"B{r}"].value), pham_vi=w[f"D{r}"].value, noi_dung=str(w[f"E{r}"].value or ""),
            dvt=str(w[f"F{r}"].value or ""), kl_hd=w[f"G{r}"].value, don_gia=num(w[f"H{r}"].value), nhom=w[f"K{r}"].value,
            ma_ns=k["nhom_ns"].get(w[f"K{r}"].value, "")))
    ma_ns_of = {(m, d["stt"]): d["ma_ns"] for m, ds in k["dong"].items() for d in ds}
    w = wb["N9_TT_DoiTac"]; k["sheet_rows"]["N9"] = 1
    for r in range(2, w.max_row + 1):
        m = w[f"A{r}"].value
        if not m: continue
        k["sheet_rows"]["N9"] = r
        loai, st = w[f"D{r}"].value, str(w[f"E{r}"].value) if w[f"E{r}"].value is not None else ""
        kl, dg, tien = w[f"H{r}"].value, num(w[f"I{r}"].value), num(w[f"J{r}"].value)
        gt = kl * dg if isinstance(kl, (int, float)) else tien
        if isinstance(w[f"B{r}"].value, (int, float)) and w[f"B{r}"].value:   # đợt cuối: tách HSTT chính (KL) và tạm ứng giữa kỳ
            key = "dot_cuoi" if loai in ("THUC_HIEN", "DIEU_CHINH", "HOAN_UNG") else "dot_cuoi_tu"
            k[key][m] = max(k[key].get(m, 0), int(w[f"B{r}"].value))
        if loai in ("THUC_HIEN", "DIEU_CHINH"):
            if isinstance(kl, (int, float)): k["lk_kl"][(m, st)] += kl
            k["lk_tien"][m] += gt; k["th_ns"][ma_ns_of.get((m, st), "")] += gt; k["lk_tien_dong"][(m, st)] += gt
        if loai in ("TAM_UNG", "HOAN_UNG"): k["tu_treo"][m] += tien
    w = wb["N2_NganSach"]
    for r in range(2, w.max_row + 1):
        if w[f"D{r}"].value:
            kl, dg = w[f"J{r}"].value, w[f"K{r}"].value
            k["ns"][w[f"D{r}"].value] += kl * dg if isinstance(kl, (int, float)) and isinstance(dg, (int, float)) else num(w[f"L{r}"].value)
    import cdt
    k["cdt"] = cdt.doc_ben_cdt(wb, k["nhom_ns"])
    wb.close()
    return k

def _chuan_hd(s): return re.sub(r"[^A-Z0-9]", "", na(s).replace("SO ", "", 1) if na(s).startswith("SO ") else na(s))
def _chuan_ten(s): return set(re.sub(r"[^A-Z0-9 ]", " ", na(s)).split()) - {"CONG", "TY", "TNHH", "CP", "CO", "PHAN", "TO", "DOI", "SX", "TM", "MTV", "DV", "XD"}

def phan_loai(hs, k):
    """HSTT → hợp đồng trong khung. Không chắc thì báo, không tự đoán."""
    cv, co = hs["cover"], []
    theo_so = next((m for m, h in k["hd"].items() if cv["so_hd"] and _chuan_hd(h["so_hd"]) == _chuan_hd(cv["so_hd"])), None)
    ten = _chuan_ten(cv["ten_don_vi"] or "")
    diem = []
    for m, h in k["hd"].items():
        t = _chuan_ten(k["doi_tac"].get(h["ma_doi_tac"], ("",))[0])
        if t and ten: diem.append((len(t & ten) / len(t), m))
    theo_ten = max(diem)[1] if diem and max(diem)[0] >= 0.75 else None
    ma_hd = theo_ten or theo_so
    if theo_so and theo_ten and theo_so != theo_ten:
        co.append(("CHAN", "Hồ sơ", cv["o"].get("so_hd", "COVER"), cv["so_hd"], k["hd"][theo_ten]["so_hd"],
                   f"Số HĐ trên COVER là HĐ của đơn vị khác ({theo_so}) — tên đơn vị khớp {theo_ten}"))
    elif theo_ten and not theo_so:
        co.append(("LUU_Y", "Hồ sơ", cv["o"].get("so_hd", "COVER"), cv["so_hd"], k["hd"][theo_ten]["so_hd"], "Số HĐ trên COVER không khớp số HĐ trong khung"))
    if not ma_hd:
        co.append(("CHAN", "Hồ sơ", cv["o"].get("ten_don_vi", "COVER"), cv["ten_don_vi"], "—", "Đơn vị CHƯA có hợp đồng trong khung — không ghi sổ (quy tắc: chưa có HĐ thì không nạp)"))
        return dict(ma_hd=None, loai_doi_tac=None, loai_hs=None), co
    loai_dt = k["doi_tac"].get(k["hd"][ma_hd]["ma_doi_tac"], ("", ""))[1]
    kn = sum(l["tt_kn"] for l in hs["lines"])
    loai_hs = "TAM_UNG" if abs(kn) < 1 and hs["du_tru"] > k["tu_treo"][ma_hd] + 1 else ("QUYET_TOAN" if "QUYET TOAN" in na(hs["file"]) else "THANH_TOAN")
    return dict(ma_hd=ma_hd, loai_doi_tac=loai_dt, loai_hs=loai_hs), co

def khop_dong(l, ds_dong):
    """Dòng HSTT → dòng HĐ trong khung: cùng nội dung (+ STT, + ĐG khi trùng). None = phát sinh chưa có trong HĐ."""
    c = [d for d in ds_dong if na(d["noi_dung"]) == na(l["ds"])]
    if len(c) > 1: c = [d for d in c if d["stt"] == l["stt"]] or [d for d in c if abs(d["don_gia"] - l["dg"]) < 0.5] or c[:1]
    return c[0] if c else None

def kiem(hs, k, van_tay_da_co=()):
    if hs.get("loai") == "CDT":
        import cdt; return cdt.kiem_cdt(hs, k, van_tay_da_co)
    pl, co = phan_loai(hs, k)
    S = hs["sheet"]; cv = hs["cover"]; t = hs["tong"] or (0, 0, 0)
    add = lambda muc, lop, vt, a, b, mt: co.append((muc, lop, vt, a, b, mt))
    if hs.get("van_tay") in van_tay_da_co: add("CHAN", "Hồ sơ", "file", hs["van_tay"][:12], "đã nạp", "File trùng nội dung với hồ sơ đã nạp trước (dù tên file khác)")
    if cv["ngay"] is None: add("LUU_Y", "Hồ sơ", "COVER", "—", "—", "Không đọc được ngày HSTT (COVER/BIA)")
    m_file = re.search(r"DOT\s*0?(\d+)", na(hs["file"].split("\\")[-1]))
    if m_file and cv["dot"] and int(m_file.group(1)) != cv["dot"]: add("LUU_Y", "Hồ sơ", cv["o"].get("dot", "COVER"), cv["dot"], int(m_file.group(1)), "Số đợt trên COVER khác số đợt trong tên file")
    # lớp 3 — số học (không cần HĐ)
    kn = sum(l["tt_kn"] for l in hs["lines"])
    if abs(kn - t[1]) > NGUONG: add("CHAN", "Số học", f"{S} dòng TỔNG", round(kn), round(t[1]), "Σ các dòng kỳ này ≠ dòng TỔNG kỳ này")
    for l in hs["lines"]:
        if abs(l["kl_kn"] * l["dg"] - l["tt_kn"]) > NGUONG: add("CHAN", "Số học", f"{S}!dòng {l['dong']}", round(l["tt_kn"]), round(l["kl_kn"] * l["dg"]), f"Tiền ≠ KL × ĐG: {l['ds'][:40]}")
        if abs(l["kl_kt"] + l["kl_kn"] - l["kl_lk"]) > 1e-6: add("CHAN", "Số học", f"{S}!dòng {l['dong']}", l["kl_lk"], l["kl_kt"] + l["kl_kn"], f"KL lũy kế ≠ kỳ trước + kỳ này: {l['ds'][:40]}")
    ma = pl["ma_hd"]; tom = dict(ky_nay=t[1], luy_ke=t[2], ky_truoc=t[0], so_dong=len(hs["lines"]), don_vi=cv["ten_don_vi"], so_hd=cv["so_hd"], dot=cv["dot"],
                                ngay=cv["ngay"].isoformat() if cv["ngay"] else None, vat=hs["vat"], du_tru_tam_ung=hs["du_tru"])
    if not ma: return dict(phan_loai=pl, co=_dang(co), tom_tat=tom, khop=[])
    h = k["hd"][ma]; ds_dong = k["dong"].get(ma, [])
    # lớp 1 — đợt
    dc = k["dot_cuoi"].get(ma, 0)
    dtu = k["dot_cuoi_tu"].get(ma, 0)
    if cv["dot"] is not None and cv["dot"] <= dc:                 # MỌI loại hồ sơ (kể cả tạm ứng): đợt ≤ đợt đã ghi ⇒ hồ sơ CŨ
        add("CHAN", "Hồ sơ", cv["o"].get("dot", "COVER"), cv["dot"], dc, f"Đợt {cv['dot']} cũ hơn/đã ghi sổ (khung đã ghi tới đợt {dc}) — chống trả trùng, chống ghi lùi lũy kế")
    elif pl["loai_hs"] == "TAM_UNG" and cv["dot"] is not None and cv["dot"] <= dtu:
        add("CHAN", "Hồ sơ", cv["o"].get("dot", "COVER"), cv["dot"], dtu, f"Tạm ứng đợt {cv['dot']} đã ghi sổ")
    elif cv["dot"] is not None and cv["dot"] > dc + 1: add("LUU_Y", "Hồ sơ", cv["o"].get("dot", "COVER"), cv["dot"], dc + 1, "Thiếu đợt ở giữa so với khung")
    # lớp 2 — theo HĐ · lớp 4 — theo đợt trước
    khop = []
    for l in hs["lines"]:
        d = khop_dong(l, ds_dong); khop.append((l, d))
        vt = f"{S}!dòng {l['dong']}"
        if d is None:
            if abs(l["kl_lk"]) > 0 or abs(l["tt_lk"]) > 0:
                add("LUU_Y", "Theo HĐ", vt, l["ds"][:40], "—", "Dòng chưa có trong HĐ — ghi sổ sẽ thêm dòng PHÁT SINH (N7), cần phụ lục")
            continue
        if d["don_gia"] > 0 and abs(l["dg"] - d["don_gia"]) > 0.5 and abs(l["kl_kn"]) > 0:
            add("LUU_Y", "Theo HĐ", vt, l["dg"], d["don_gia"], f"ĐG khác HĐ ⇒ thiếu báo giá / PLHĐ: {l['ds'][:40]}")
        if isinstance(d["kl_hd"], (int, float)) and d["kl_hd"] > 0 and l["kl_lk"] > d["kl_hd"] * 1.0001 + 0.001:
            add("LUU_Y", "Theo HĐ", vt, l["kl_lk"], d["kl_hd"], f"KL lũy kế vượt KL HĐ: {l['ds'][:40]}")
        if l["dvt"] and d["dvt"] and na(l["dvt"]) != na(d["dvt"]): add("LUU_Y", "Theo HĐ", vt, l["dvt"], d["dvt"], f"ĐVT khác HĐ: {l['ds'][:40]}")
        lk_khung = k["lk_kl"].get((ma, d["stt"]), 0.0)
        if abs(l["kl_kt"] - lk_khung) > 1e-6 and not any(x is not l and khop_dong(x, ds_dong) is d for x in hs["lines"]):
            add("LUU_Y", "Đợt trước", vt, round(l["kl_kt"], 3), round(lk_khung, 3), f"KL kỳ trước ≠ lũy kế đã ghi sổ ⇒ sẽ ghi ĐIỀU CHỈNH: {l['ds'][:40]}")
    if abs(t[0] - k["lk_tien"][ma]) > NGUONG: add("LUU_Y", "Đợt trước", f"{S} dòng TỔNG", round(t[0]), round(k["lk_tien"][ma]), "Giá trị kỳ trước ≠ lũy kế đã ghi sổ trong khung")
    if h["pct_tt_dot"] is not None and hs["gl"] is not None and t[2] > 0:
        pct_gl = -hs["gl"] / (t[2] * (1 + num(hs["vat"])))
        if abs(pct_gl - (1 - num(h["pct_tt_dot"]))) > 0.005: add("LUU_Y", "Theo HĐ", f"{S} dòng (GL)", f"{pct_gl:.1%}", f"{1 - num(h['pct_tt_dot']):.1%}", "% giữ lại khác HĐ")
    if hs["du_tru"] > 0 and h["pct_tu"] == 0: add("LUU_Y", "Theo HĐ", f"{S} dòng (TU)", round(hs["du_tru"]), 0, "Tạm ứng dù HĐ ghi 'không áp dụng'")
    # ngân sách theo mã NS sau khi ghi đợt này
    them = defaultdict(float)
    for l, d in khop:
        if d: them[d["ma_ns"]] += l["tt_kn"]
    for mn, v in them.items():
        if mn and k["ns"].get(mn, 0) > 0 and k["th_ns"][mn] + v > k["ns"][mn]:
            add("LUU_Y", "Ngân sách", mn, round(k["th_ns"][mn] + v), round(k["ns"][mn]), f"Sau đợt này thực hiện VƯỢT ngân sách mã {mn}")
    return dict(phan_loai=pl, co=_dang(co), tom_tat=tom, khop=[(l, d["stt"] if d else None) for l, d in khop])

def _dang(co):
    return [dict(muc=m, lop=lop, vi_tri=vt, hstt=a, doi_chieu=b, mo_ta=mt) for m, lop, vt, a, b, mt in co]
