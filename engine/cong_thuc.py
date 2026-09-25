"""CÔNG THỨC DÒNG của khung CCM v3 — sao y tao_khung_v3.py (ct_sheet / tt_sheet), để ghi được vào khung TRỐNG (không có dòng mẫu).
Đã so khớp TỪNG KÝ TỰ với khung sinh bởi tao_khung_v3.py (xem kiem_cong_thuc() ở cuối)."""
N = 5000
NH = "'N1_DanhMuc'!$Q$3:$T$400"
HD_COLS = ["ma_hd", "loai_ghi", "ma_doi_tac", "so_hd", "ngay_ky", "noi_dung", "dang_hd", "gia_tri_truoc_vat", "vat_pct", "pct_tam_ung", "pct_tt_dot",
           "pct_tt_quyet_toan", "han_tt_ngay", "don_vi_han", "han_qt_ngay", "han_tra_gl_ngay", "bao_hanh_thang", "ngay_nt_hoan_thanh", "gia_tri_quyet_toan",
           "ngay_quyet_toan", "trang_thai", "nguon", "ghi_chu", "ho_so_thieu", "hs_qt_bb_nghiem_thu", "hs_qt_bang_kl", "hs_qt_thanh_ly", "hs_qt_bl_bao_hanh"]
COT_HD = {"N4_HD_CDT": HD_COLS, "N6_HD_DoiTac": HD_COLS[:3] + ["ma_goi"] + HD_COLS[3:]}
CT_COLS = ["ma_hd", "stt", "khoa", "pham_vi", "noi_dung", "dvt", "kl_hd", "don_gia", "thanh_tien", "ma_cv", "nhom", "ma_ns", "kl_luy_ke", "kiem_tra", "nguon"]
TT_COLS = ["ma_hd", "dot", "ngay", "loai", "stt_dong", "noi_dung", "dvt", "kl_ky_nay", "don_gia", "so_tien", "gia_tri_ky", "nhom", "ma_ns",
           "tien_thanh_toan", "ngay_den_han", "ngay_da_tra", "kiem_tra", "nguon"]
CAP = {"N5_BOQ_CDT": ("N8_TT_CDT", "N4_HD_CDT"), "N7_HD_DoiTac_ChiTiet": ("N9_TT_DoiTac", "N6_HD_DoiTac")}      # sheet dòng HĐ → (sheet TT, sheet HĐ)
CAP_TT = {"N8_TT_CDT": ("N5_BOQ_CDT", "N4_HD_CDT"), "N9_TT_DoiTac": ("N7_HD_DoiTac_ChiTiet", "N6_HD_DoiTac")}
def L(i):
    s = ""
    while i: i, r = divmod(i - 1, 26); s = chr(65 + r) + s
    return s
def cot(cols): return {k: L(i) for i, k in enumerate(cols, 1)}
def R(sheet, cols, k): c = cot(cols)[k]; return f"'{sheet}'!${c}$2:${c}${N}"
def VL(hd_sheet, key, k):
    cs = COT_HD[hd_sheet]; c = cot(cs); return f"VLOOKUP({key},'{hd_sheet}'!${c['ma_hd']}:${c[k]},{cs.index(k) - cs.index('ma_hd') + 1},0)"

def ct_cong_thuc(ct_sheet, r):
    """Công thức các cột tính của 1 dòng N5/N7: {chữ cột: công thức}."""
    tt = CAP[ct_sheet][0]; c = cot(CT_COLS)
    return {c["khoa"]: f'={c["ma_hd"]}{r}&"|"&{c["stt"]}{r}',
            c["thanh_tien"]: f'=IF({c["kl_hd"]}{r}="","",{c["kl_hd"]}{r}*{c["don_gia"]}{r})',
            c["ma_ns"]: f'=IFERROR(VLOOKUP({c["nhom"]}{r},{NH},4,0),"")',
            c["kl_luy_ke"]: f'=SUMIFS({R(tt, TT_COLS, "kl_ky_nay")},{R(tt, TT_COLS, "ma_hd")},{c["ma_hd"]}{r},{R(tt, TT_COLS, "stt_dong")},{c["stt"]}{r})',
            c["kiem_tra"]: (f'=IF({c["nhom"]}{r}="","CHƯA GÁN NHÓM",IF(ISNA(VLOOKUP({c["nhom"]}{r},{NH},3,0)),"NHÓM LẠ",'
                            f'IF({c["dvt"]}{r}<>VLOOKUP({c["nhom"]}{r},{NH},3,0),"LỆCH ĐVT",'
                            f'IF(AND(ISNUMBER({c["kl_hd"]}{r}),ABS({c["kl_luy_ke"]}{r})>ABS({c["kl_hd"]}{r})*1.0001+0.001),"VƯỢT KL HĐ",'
                            f'IF(LEFT({c["nhom"]}{r},3)="BG_","BÙ GIÁ","")))))')}

def tt_cong_thuc(tt_sheet, r, co_noi_dung=False):
    """Công thức các cột tính của 1 dòng N8/N9 (khi dòng có nội dung riêng thì F/G là giá trị, không phải công thức)."""
    ct, hd = CAP_TT[tt_sheet]; c = cot(TT_COLS)
    key = f'{c["ma_hd"]}{r}&"|"&{c["stt_dong"]}{r}'; isth = f'OR({c["loai"]}{r}="THUC_HIEN",{c["loai"]}{r}="DIEU_CHINH")'
    v = lambda k: VL(hd, f'{c["ma_hd"]}{r}', k)
    f = {c["gia_tri_ky"]: f'=IF(AND({isth},{c["kl_ky_nay"]}{r}<>""),{c["kl_ky_nay"]}{r}*{c["don_gia"]}{r},{c["so_tien"]}{r})',
         c["nhom"]: f"=IFERROR(VLOOKUP({key},'{ct}'!$C:$K,9,0),\"\")", c["ma_ns"]: f"=IFERROR(VLOOKUP({key},'{ct}'!$C:$L,10,0),\"\")",
         c["tien_thanh_toan"]: f'=IF({isth},{c["gia_tri_ky"]}{r}*(1+{v("vat_pct")})*{v("pct_tt_dot")},{c["gia_tri_ky"]}{r})',
         c["ngay_den_han"]: f'=IF({c["ngay"]}{r}="","",IF({v("don_vi_han")}="LV",WORKDAY({c["ngay"]}{r},{v("han_tt_ngay")}),{c["ngay"]}{r}+{v("han_tt_ngay")}))',
         c["kiem_tra"]: (f'=IF({c["loai"]}{r}<>"THUC_HIEN","",IF(ISNA(VLOOKUP({key},\'{ct}\'!$C:$C,1,0)),"NGOÀI DS DÒNG HĐ",'
                         f'IF(VLOOKUP({key},\'{ct}\'!$C:$H,6,0)=0,"",IF(ABS({c["don_gia"]}{r}-VLOOKUP({key},\'{ct}\'!$C:$H,6,0))>0.5,"ĐG≠HĐ",""))))')}
    if not co_noi_dung:
        f[c["noi_dung"]] = f"=IFERROR(VLOOKUP({key},'{ct}'!$C:$E,3,0),\"\")"; f[c["dvt"]] = f"=IFERROR(VLOOKUP({key},'{ct}'!$C:$F,4,0),\"\")"
    return f

def _bo_nhay(x):                                                     # Excel bỏ dấu nháy quanh tên sheet không có khoảng trắng khi lưu
    import re; return re.sub(r"'([A-Za-z0-9_]+)'!", lambda m: m.group(1) + "!", str(x))
def kiem_cong_thuc(khung_mau):
    """So công thức sinh ra với khung mẫu có dữ liệu (tao_khung_v3.py). Trả danh sách ô lệch."""
    import openpyxl
    wb = openpyxl.load_workbook(khung_mau); lech = []
    for sh, fn in (("N5_BOQ_CDT", ct_cong_thuc), ("N7_HD_DoiTac_ChiTiet", ct_cong_thuc), ("N8_TT_CDT", tt_cong_thuc), ("N9_TT_DoiTac", tt_cong_thuc)):
        ws = wb[sh]
        for r in (2, 3, ws.max_row):
            if not ws[f"A{r}"].value: continue
            co_nd = sh in CAP_TT and not str(ws[f"F{r}"].value or "").startswith("=")
            for col, cth in (fn(sh, r, co_nd) if sh in CAP_TT else fn(sh, r)).items():
                if _bo_nhay(ws[f"{col}{r}"].value) != _bo_nhay(cth): lech.append((sh, f"{col}{r}", ws[f"{col}{r}"].value, cth))
    return lech
