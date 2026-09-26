"""ENGINE SOÁT HSTT — chạy trên máy anh:  python app.py      (mặc định http://127.0.0.1:8765)
Trang (GitHub Pages / file local) gọi vào đây. Dữ liệu CHỈ nằm trên máy: D:\\QLCP_HD\\WEBAPP_SOAT_HSTT_DATA\\<dự án>\\
  GET  /ping                      · GET /du-an                 · GET /ho-so?du_an=X
  POST /nap    {du_an, ten, b64}  → lưu file, phân loại, tự kiểm 4 lớp, kế hoạch ghi sổ
  POST /duyet  {du_an, id, hanh_dong: DONG_Y | YEU_CAU_SUA | TRA_DOI, ly_do}
  GET  /bao-cao?du_an=X           → R0 tổng quan + 90_Check (đọc từ file khung)
  GET  /danh-muc?du_an=X · GET /ho-so-nen?du_an=X   → danh sách chọn (đối tác, HĐ, gói) · sổ hồ sơ nền
  POST /nap-nen {du_an, ten, b64, loai, ma_dt, loai_dt, goi, ghi_chu} → lưu BoQ/NS/gói/báo giá/HĐ/QT đúng folder
  POST /doan-hstt {du_an, ten} · POST /dinh-kem {du_an, id, ten, b64} → PDF bản ký gắn vào HSTT Excel cùng HĐ + đợt
  GET  /du-lieu?du_an=X           → ĐẦU RA: hợp đồng · bill · báo cáo tài chính · đối tác · dòng tiền (đọc R1…R7 của file khung)"""
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import json, os, sys, base64, datetime as dt, threading, traceback
from urllib.parse import urlparse, parse_qs
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import doc_hstt as D, kiem as K, ghi_so as G, bao_cao as BC, cay as C, nen as NEN, nhap_khung as NK

PORT = 8765
# VỊ TRÍ DỮ LIỆU — chỉ ghi ở máy này (engine/cau_hinh_may.json, không lên git): {"DATA": "<thư mục dữ liệu>"}.
# Đổi công ty / đổi ổ / đổi máy ⇒ chỉ sửa 1 dòng này. Mọi đường dẫn bên trong DATA đều lưu TƯƠNG ĐỐI nên không phải sửa gì thêm.
MAY = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cau_hinh_may.json")
MAY_CFG = json.load(open(MAY, encoding="utf-8")) if os.path.exists(MAY) else {}
DATA = MAY_CFG.get("DATA", r"D:\QLCP_HD\WEBAPP_SOAT_HSTT_DATA")
CTY = MAY_CFG.get("CONG_TY", "(chưa khai báo công ty — thêm CONG_TY vào engine/cau_hinh_may.json)")   # để AI biết mình là bên nào trong HĐ
def tuyet_doi(x): return os.path.normpath(os.path.join(DATA, x)) if x and not os.path.isabs(x) else x
def tuong_doi(x):
    try: return os.path.relpath(x, DATA) if x and os.path.isabs(x) and os.path.normcase(x).startswith(os.path.normcase(DATA)) else x
    except ValueError: return x                                     # khác ổ đĩa ⇒ giữ nguyên
DUONG = ("file", "luu_tru")                                           # các trường đường dẫn trong so_nap.json
def _cfg(ten):                                        # cấu hình ở _CAU_HINH\ (cây mới); còn file ở gốc (cây cũ) thì vẫn đọc được
    moi = os.path.join(DATA, "_CAU_HINH", ten); return moi if os.path.exists(moi) or not os.path.exists(os.path.join(DATA, ten)) else os.path.join(DATA, ten)
CAU_HINH = _cfg("du_an.json")          # {"DU_AN_A": {"ten": "...", "khung": "<đường dẫn file khung .xlsx>"}}
TRANG_TXT = _cfg("trang.txt")          # 1 dòng: https://<tài-khoản>.github.io — trang khác KHÔNG gọi được engine
TRANG = open(TRANG_TXT, encoding="utf-8").read().strip().rstrip("/") if os.path.exists(TRANG_TXT) else ""
KHOA = threading.Lock()                               # 1 lần ghi sổ tại 1 thời điểm

def _json(o):
    if isinstance(o, (dt.date, dt.datetime)): return o.isoformat()
    if isinstance(o, dt.time): return None if o == dt.time(0) else o.isoformat()   # ô ngày TRỐNG qua công thức = 0 ⇒ openpyxl đọc thành 00:00 ⇒ coi là chưa khai báo
    raise TypeError(type(o))
def du_an():
    d = json.load(open(CAU_HINH, encoding="utf-8")) if os.path.exists(CAU_HINH) else {}
    for v in d.values(): v["khung"] = tuyet_doi(v.get("khung"))
    return d
def he_thong(da): return os.path.join(DATA, da, "_HE_THONG")             # engine quản lý — không sửa tay
def so_nap_path(da): return os.path.join(he_thong(da), "so_nap.json")
def so_nap(da):
    p = so_nap_path(da); s = json.load(open(p, encoding="utf-8")) if os.path.exists(p) else {}
    for r in s.values():
        for f in DUONG:
            if r.get(f): r[f] = tuyet_doi(r[f])
        for x in r.get("dinh_kem", []): x["file"] = tuyet_doi(x["file"])
    return s
def luu_so(da, s):
    os.makedirs(he_thong(da), exist_ok=True)
    s = {k: dict(v, **{f: tuong_doi(v[f]) for f in DUONG if v.get(f)},
                 **({"dinh_kem": [dict(x, file=tuong_doi(x["file"])) for x in v["dinh_kem"]]} if v.get("dinh_kem") else {})) for k, v in s.items()}
    json.dump(s, open(so_nap_path(da), "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=_json)

def xu_ly_nap(b):
    da, ten = b["du_an"], os.path.basename(b["ten"])
    if not ten.lower().endswith((".xlsx", ".xlsm")):              # PDF / Word / ảnh: chưa có luồng đọc ⇒ báo rõ, KHÔNG lưu file mồ côi
        raise ValueError(f"'{ten}' không phải HSTT dạng Excel. Hiện webapp chỉ nạp HSTT (.xlsx). Hợp đồng, BoQ, ngân sách, báo giá (PDF/Excel) "
                         "sẽ nạp được ở chức năng NẠP HỒ SƠ NỀN — tạm thời anh bỏ file vào đúng folder trên Drive.")
    cfg = du_an()[da]; thu_muc = os.path.join(he_thong(da), "nap"); os.makedirs(thu_muc, exist_ok=True)
    raw = base64.b64decode(b["b64"]); tam = os.path.join(thu_muc, "_tam_" + ten)
    open(tam, "wb").write(raw); vt = D.van_tay(tam)
    s = so_nap(da); i = vt[:12]
    cu = {v["van_tay"] for k_, v in s.items() if k_ != i and v["trang_thai"] != "TRA_DOI"}     # trùng nội dung với HỒ SƠ KHÁC
    dich = os.path.join(thu_muc, f"{i}_{ten}")
    if os.path.exists(dich): os.remove(tam)                        # nạp lại đúng file cũ: giữ bản gốc chỉ đọc
    else: os.replace(tam, dich); os.chmod(dich, 0o444)
    hs = D.doc_file(dich); hs["van_tay"] = vt
    k = K.doc_khung(cfg["khung"]); k["tu_khoa"] = cfg.get("tu_khoa", []); kq = K.kiem(hs, k, cu)
    if hs.get("mau") == "CONG_NHAT" and hs.get("lk_cover") is not None and abs(hs["tong"][2] - hs["lk_cover"]) > 1:   # công nhật: Σ ngày công = LK trên COVER
        kq["co"].insert(0, dict(muc="CHAN", lop="Số học", vi_tri="COVER", hstt=hs["tong"][2], doi_chieu=hs["lk_cover"],
                                mo_ta=f"Σ ngày công {hs['tong'][2]:,.0f} ≠ giá trị thực hiện lũy kế trên COVER {hs['lk_cover']:,.0f} — bảng KL thiếu/thừa dòng"))
    for da2 in du_an():                                            # HĐ nguyên tắc NCC dùng chung nhiều dự án ⇒ cùng 1 file không được nạp ở 2 dự án
        if da2 == da: continue
        r2 = next((v for v in so_nap(da2).values() if v["van_tay"] == vt and v["trang_thai"] != "TRA_DOI"), None)
        if r2: kq["co"].insert(0, dict(muc="CHAN", lop="Hồ sơ", vi_tri="file", hstt=vt[:12], doi_chieu=da2,
                                        mo_ta=f"Cùng file này đã nạp ở dự án {da2} ({r2['trang_thai']}) — không ghi 2 dự án"))
    kh = G.ke_hoach(hs, kq, k) if kq["phan_loai"]["ma_hd"] else None
    cu_rec = s.get(i)
    if cu_rec and cu_rec["trang_thai"] == "DA_GHI_SO":           # đã ghi sổ ⇒ không cho ghi lần 2
        cu_rec["co"] = [dict(muc="CHAN", lop="Hồ sơ", vi_tri="file", hstt=i, doi_chieu="đã ghi sổ", mo_ta="Hồ sơ này ĐÃ GHI SỔ lúc " + str(cu_rec.get("luc_duyet")))]
        return cu_rec
    rec = dict(id=i, van_tay=vt, ten=ten, file=dich, luc=dt.datetime.now().isoformat(timespec="seconds"),
               trang_thai="CHO_DUYET",                                   # nạp lại (kể cả file đã TRẢ ĐỘI / YÊU CẦU SỬA) ⇒ soát lại từ đầu, chờ duyệt
               phan_loai=kq["phan_loai"], tom_tat=kq["tom_tat"], co=kq["co"], ke_hoach=kh, ly_do=None, ket_qua=None,
               lich_su=((cu_rec or {}).get("lich_su") or []) + ([dict(trang_thai=cu_rec["trang_thai"], ly_do=cu_rec.get("ly_do"), luc=cu_rec.get("luc_duyet") or cu_rec.get("luc"))]
                                                        if cu_rec and cu_rec["trang_thai"] != "CHO_DUYET" else []))
    if cu_rec and cu_rec.get("hd_tam"): rec["hd_tam"] = cu_rec["hd_tam"]
    s[i] = rec; luu_so(da, s)
    return rec

def xu_ly_duyet(b):
    da, i, hd = b["du_an"], b["id"], b["hanh_dong"]
    s = so_nap(da); rec = s[i]
    if rec["trang_thai"] == "DA_GHI_SO": return dict(ok=False, ly_do="Hồ sơ đã ghi sổ trước đó")
    if hd == "DONG_Y":
        if any(c["muc"] == "CHAN" for c in rec["co"]): return dict(ok=False, ly_do="Còn cờ CHẶN — nút Đồng ý bị khoá")
        with KHOA:
            cfg = du_an()[da]
            hs = D.doc_file(rec["file"]); hs["van_tay"] = rec["van_tay"]
            k = K.doc_khung(cfg["khung"]); k["tu_khoa"] = cfg.get("tu_khoa", []); kq = K.kiem(hs, k)                  # kiểm lại ngay trước khi ghi (khung có thể đã đổi)
            if kq["phan_loai"].get("ma_hd") == "HD-CDT" and (hs.get("tien") or {}).get("khau_tru") and "HD-CDT-KT" not in k["hd"]:   # CĐT khấu trừ ⇒ cần HĐ bên CHI
                t = NK.tao_hd_cdt_kt(cfg["khung"], hs.get("vat") or 0, os.path.join(os.path.dirname(cfg["khung"]), "_backup"), nguon=f"webapp · {rec['ten'][:40]}")
                if not t["ok"]: return dict(ok=False, ly_do="Không tạo được HĐ khấu trừ HD-CDT-KT: " + t["ly_do"])
                k = K.doc_khung(cfg["khung"]); k["tu_khoa"] = cfg.get("tu_khoa", []); kq = K.kiem(hs, k)
            if any(c["muc"] == "CHAN" for c in kq["co"]): return dict(ok=False, ly_do="Kiểm lại trước khi ghi phát sinh cờ CHẶN", co=kq["co"])
            kh = G.ke_hoach(hs, kq, k)
            kq_ghi = G.ghi(cfg["khung"], kh, rec["ten"], os.path.join(os.path.dirname(cfg["khung"]), "_backup"))
        rec["ket_qua"] = kq_ghi
        if kq_ghi["ok"]:
            rec["trang_thai"] = "DA_GHI_SO"
            try: rec["luu_tru"] = C.luu_tru(DATA, da, rec, cfg["khung"]); NEN.xep_dinh_kem(rec)   # xếp bản gốc + PDF bản ký vào …\HSTT\Dxx_YYYYMMDD\
            except Exception as e: rec["luu_tru_loi"] = str(e)
    else:
        rec["trang_thai"] = hd; rec["ly_do"] = b.get("ly_do") or ""
    rec["luc_duyet"] = dt.datetime.now().isoformat(timespec="seconds"); s[i] = rec; luu_so(da, s)
    return dict(ok=rec["trang_thai"] != "CHO_DUYET", ho_so=rec)

def bao_cao(da):
    import openpyxl
    wb = openpyxl.load_workbook(du_an()[da]["khung"], data_only=True, read_only=True)
    r0 = [(r[1], r[2]) for r in wb["R0_TongQuan"].iter_rows(min_row=4, max_row=30, max_col=4, values_only=True) if r[1]]
    ck = [dict(so=r[0], ten=r[1], trai=r[2], phai=r[3], ket_luan=r[5]) for r in wb["90_Check"].iter_rows(min_row=2, max_row=60, max_col=7, values_only=True) if r[1]]
    wb.close()
    return dict(tong_quan=r0, kiem=ck, ghi_chu="Số đọc từ lần Excel tính gần nhất (sau mỗi lần ghi sổ).")

_DL = {}                                              # cache theo (đường dẫn, mtime) — file đổi (ghi sổ) thì đọc lại
def du_lieu(da):
    p = du_an()[da]["khung"]; m = os.path.getmtime(p)
    if _DL.get(p, (None,))[0] != m: _DL[p] = (m, BC.doc(p, so_nap(da)))
    return dict(_DL[p][1], cap_nhat=dt.datetime.fromtimestamp(m).strftime("%d/%m/%Y %H:%M"))

def xu_ly_nap_nen(b):
    da = b["du_an"]; cfg = du_an()[da]
    return NEN.nap_nen(DATA, da, cfg["khung"], b["ten"], base64.b64decode(b["b64"]), b.get("loai") or None, b.get("ma_dt"), b.get("loai_dt"), b.get("goi"), b.get("ghi_chu", ""))
def xu_ly_nhap_khung(b):
    """Anh duyệt hồ sơ nền ⇒ ghi vào khung (HĐ CĐT / đối tác). Kiểm lại cờ CHẶN ngay trước khi ghi."""
    da, i = b["du_an"], b["id"]; cfg = du_an()[da]
    with KHOA:
        rec = NEN.doc_so(DATA, da)[i]
        if rec.get("trang_thai") != "CHO_DUYET": raise ValueError("Hồ sơ này không ở trạng thái chờ duyệt")
        if any(c["muc"] == "CHAN" for c in rec.get("co", [])): raise ValueError("Còn cờ CHẶN — chưa ghi khung được")
        if rec["loai"] == "NGAN_SACH":                                   # ngân sách R00 theo mẫu nội bộ — bộ đọc tất định
            import ns_r00, datetime as _d
            f = os.path.join(DATA, rec["duong_dan"])
            if not ns_r00.la_mau_r00(f): raise ValueError("File ngân sách không theo mẫu R00 nội bộ (cần sheet BCTC + 01. PhanTichBOQ)")
            hd_cdt = [h for h in NEN.danh_muc(cfg["khung"])["hop_dong"] if h["ben"] == "CĐT"]
            kq = ns_r00.ghi_r00(cfg["khung"], f, _d.date.today(), os.path.join(os.path.dirname(cfg["khung"]), "_backup"))
            NEN.sua_rec(DATA, da, i, ket_qua_khung=kq, **({"trang_thai": "DA_NHAP", "da_nhap_khung": True} if kq["ok"] else {}))
            if not kq["ok"]: raise ValueError(kq["ly_do"])
            return kq
        if rec["loai"] not in ("HD_CDT", "HD_DOI_TAC"): raise ValueError("Ghi khung cho BoQ / gói thầu đang làm — file đã lưu đúng folder")
        if rec.get("la_phu_luc") or str((rec.get("ai") or {}).get("loai") or "").startswith("PLHD"):     # PHỤ LỤC: không bao giờ ghi thành HĐ gốc
            kq = NK.ghi_phu_luc(cfg["khung"], da, rec, os.path.join(os.path.dirname(cfg["khung"]), "_backup"))
            NEN.sua_rec(DATA, da, i, ket_qua_khung=kq, **({"trang_thai": "DA_NHAP", "da_nhap_khung": True} if kq["ok"] else {}))
            if not kq["ok"]: raise ValueError(kq["ly_do"])
            return kq
        kq = NK.ghi_hd(cfg["khung"], da, rec, b.get("sua") or {}, os.path.join(os.path.dirname(cfg["khung"]), "_backup"))
        NEN.sua_rec(DATA, da, i, ket_qua_khung=kq, **({"trang_thai": "DA_NHAP", "da_nhap_khung": True, "ma_hd": kq["ma_hd"]} if kq["ok"] else {}))
        if kq["ok"]: C.tao_cay(DATA, da, cfg["khung"])
    if not kq["ok"]: raise ValueError(kq["ly_do"])
    return kq
LY_DO_TAM = {"CONG_NHAT": "Công nhật không có HĐ", "MUA_LE": "Mua lẻ", "HOAN_UNG_BCH": "Hoàn ứng BCH", "KHAC": "Khác"}
def xu_ly_hd_tam(b):
    """HĐ TẠM tự khai báo từ 1 HSTT chưa có HĐ (công nhật, mua lẻ, hoàn ứng BCH…): tạo đối tác + HĐ + bảng đơn giá theo dòng HSTT,
    % TT 100%, đánh dấu 'Hồ sơ còn thiếu = HĐ chính thức'. Rồi soát lại HSTT đó."""
    import re as _re
    da, i = b["du_an"], b["id"]; cfg = du_an()[da]; ly_do = LY_DO_TAM.get(b.get("ly_do"), b.get("ly_do") or "Khác"); loai_dt = b.get("loai_dt") or "DTC"
    with KHOA:
        rec = so_nap(da)[i]; hs = D.doc_file(rec["file"]); cv = hs["cover"]
        ten = _re.sub(r"\s{2,}", " ", _re.sub(r"\([^)]*\)", " ", str(cv.get("ten_don_vi") or ""))).strip(" -–,;")
        if not ten: raise ValueError("HSTT không ghi tên đơn vị — không tạo được HĐ tạm")
        dm = NEN.danh_muc(cfg["khung"]); d = NEN.khop_doi_tac(ten, dm["doi_tac"]); ma = d["ma"] if d else NEN.ma_de_xuat(ten)
        if any(h["ma_hd"] == f"HD-{ma}" for h in dm["hop_dong"]): raise ValueError(f"HD-{ma} đã có trong khung — không cần HĐ tạm")
        hom_nay = dt.date.today().strftime("%d/%m/%Y"); ngay = cv.get("ngay")
        a = dict(loai="HD_DOI_TAC", so_hd=cv.get("so_hd") or f"HĐ TẠM {ma}", ngay_ky=str(ngay)[:10] if ngay else dt.date.today().isoformat(), doi_tac_ten=ten,
                 loai_doi_tac=loai_dt, noi_dung=f"HĐ TẠM tự khai báo — {ly_do}", dang_hd="DON_GIA", gia_tri_truoc_vat=None, vat_pct=hs.get("vat") or 0,
                 pct_tam_ung=0, pct_tt_dot=1, pct_tt_quyet_toan=1, han_tt_ngay=0, don_vi_han="LICH",
                 bang=[dict(stt=l["stt"], noi_dung=l["ds"], dvt=l["dvt"], kl=None, don_gia=l["dg"], thanh_tien=None) for l in hs["lines"]],
                 nguon=f"HĐ TẠM tự khai báo từ HSTT {rec['ten'][:50]} · {hom_nay}",
                 ghi_chu=f"CHƯA CÓ HĐ CHÍNH THỨC — HĐ tạm ({ly_do}) tạo ngày {hom_nay} từ HSTT; bổ sung HĐ thật thì cập nhật lại",
                 ho_so_thieu=f"HĐ chính thức (đang dùng HĐ TẠM — {ly_do})")
        kq = NK.ghi_hd(cfg["khung"], da, dict(loai="HD_DOI_TAC", ma_doi_tac=ma, loai_doi_tac=loai_dt, ten=rec["ten"], goi=None, ai=a), {},
                       os.path.join(os.path.dirname(cfg["khung"]), "_backup"))
        if not kq["ok"]: raise ValueError(kq["ly_do"])
        C.tao_cay(DATA, da, cfg["khung"])
    r = xu_ly_nap(dict(du_an=da, ten=rec["ten"], b64=base64.b64encode(open(rec["file"], "rb").read()).decode()))       # soát lại với HĐ tạm vừa tạo
    s = so_nap(da); s[r["id"]]["hd_tam"] = dict(ma_hd=kq["ma_hd"], ly_do=ly_do, luc=dt.datetime.now().isoformat(timespec="seconds")); luu_so(da, s)
    return s[r["id"]]
def xu_ly_doan_hstt(b):
    ds = list(so_nap(b["du_an"]).values()); diem = dict((i, d) for d, i in NEN.doan_hstt(b["ten"], ds))
    return sorted([dict(id=r["id"], ten=r["ten"], ma_hd=(r.get("phan_loai") or {}).get("ma_hd"), dot=(r.get("tom_tat") or {}).get("dot"),
                        trang_thai=r["trang_thai"], so_pdf=len(r.get("dinh_kem", [])), diem=diem.get(r["id"], 0)) for r in ds], key=lambda x: -x["diem"])
def xu_ly_dinh_kem(b):
    da = b["du_an"]; s = so_nap(da)
    if b.get("id") not in s: raise ValueError("Chưa chọn HSTT Excel để gắn PDF")
    rec = NEN.dinh_kem(DATA, da, s[b["id"]], b["ten"], base64.b64decode(b["b64"])); s[b["id"]] = rec; luu_so(da, s); return rec

class H(BaseHTTPRequestHandler):
    def _cors(self):
        o = self.headers.get("Origin", "")
        if (TRANG and o == TRANG) or o.startswith("http://localhost:") or o.startswith("http://127.0.0.1:"):   # CHỈ trang của anh + chạy thử tại máy
            self.send_header("Access-Control-Allow-Origin", o or "*"); self.send_header("Vary", "Origin")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.send_header("Access-Control-Allow-Private-Network", "true")
    def _tra(self, code, obj):
        b = json.dumps(obj, ensure_ascii=False, default=_json).encode("utf-8")
        self.send_response(code); self._cors(); self.send_header("Content-Type", "application/json; charset=utf-8"); self.end_headers(); self.wfile.write(b)
    def do_OPTIONS(self): self.send_response(204); self._cors(); self.end_headers()
    def do_GET(self):
        u = urlparse(self.path); q = {k: v[0] for k, v in parse_qs(u.query).items()}
        try:
            if u.path == "/ping": return self._tra(200, dict(ok=True, engine="soat-hstt", phien_ban="1.0"))
            if u.path == "/du-an": return self._tra(200, {k: v.get("ten", k) for k, v in du_an().items()})
            if u.path == "/ho-so": return self._tra(200, sorted(so_nap(q["du_an"]).values(), key=lambda r: r["luc"], reverse=True))
            if u.path == "/bao-cao": return self._tra(200, bao_cao(q["du_an"]))
            if u.path == "/du-lieu": return self._tra(200, du_lieu(q["du_an"]))
            if u.path == "/danh-muc": return self._tra(200, NEN.danh_muc(du_an()[q["du_an"]]["khung"]))
            if u.path == "/ho-so-nen": return self._tra(200, sorted(NEN.doc_so(DATA, q["du_an"]).values(), key=lambda r: r["luc"], reverse=True))
            self._tra(404, dict(loi="không có đường dẫn này"))
        except Exception as e: traceback.print_exc(); self._tra(500, dict(loi=str(e)))
    def do_POST(self):
        try:
            b = json.loads(self.rfile.read(int(self.headers.get("Content-Length", 0))) or b"{}")
            if self.path == "/nap": return self._tra(200, xu_ly_nap(b))
            if self.path == "/duyet": return self._tra(200, xu_ly_duyet(b))
            if self.path == "/nap-nen": return self._tra(200, xu_ly_nap_nen(b))
            if self.path == "/doan-hstt": return self._tra(200, xu_ly_doan_hstt(b))
            if self.path == "/dinh-kem": return self._tra(200, xu_ly_dinh_kem(b))
            if self.path == "/nhap-khung": return self._tra(200, xu_ly_nhap_khung(b))
            if self.path == "/hd-tam": return self._tra(200, xu_ly_hd_tam(b))
            if self.path == "/doc-lai": return self._tra(200, NEN.doc_lai(DATA, b["du_an"], b["id"]))
            self._tra(404, dict(loi="không có đường dẫn này"))
        except Exception as e: traceback.print_exc(); self._tra(500, dict(loi=str(e)))
    def log_message(self, fmt, *a): sys.stderr.write(f"[{dt.datetime.now():%H:%M:%S}] {fmt % a}\n")

if __name__ == "__main__":
    for st in (sys.stdout, sys.stderr): st.reconfigure(encoding="utf-8", errors="replace")   # console Windows cp1252 không in được tiếng Việt
    os.makedirs(DATA, exist_ok=True)
    NEN.khoi_dong(DATA, du_an, CTY)                                  # luồng AI đọc hồ sơ nền (gói Claude của người dùng) + xếp lại việc dở
    for da_, v in du_an().items():                                   # bổ sung cây thư mục (idempotent) — đối tác mới có HĐ thì có folder
        try: C.tao_cay(DATA, da_, v.get("khung"))
        except Exception as e: print(f"[cây] {da_}: {e}")
    port = int(sys.argv[1]) if len(sys.argv) > 1 else PORT
    print(f"Engine soát HSTT chạy tại http://127.0.0.1:{port}  ·  dữ liệu: {DATA}")
    ThreadingHTTPServer(("127.0.0.1", port), H).serve_forever()
