"""SO GIÁ — GIAI ĐOẠN CHỌN ĐƠN VỊ THI CÔNG / NHÀ CUNG CẤP (anh Phương 29/09/2026).
- Báo giá ứng viên KHÔNG thuộc hồ sơ nền (hồ sơ nền = BoQ/HĐ CĐT, ngân sách, chủ trương) ⇒ nạp thẳng vào PHIÊN SO GIÁ.
- KHÔNG ghi khung (khung chỉ ghi khi nạp HĐ / HSTT). Kết quả chọn thầu CHỈ LƯU: sổ _HE_THONG/so_gia.json + folder phiên
  01_THIET_LAP/03_GOI_THAU/CHON_THAU/<DD.MM.YYYY>_<tên gói>/ (bản báo giá gốc + PTLN_….xlsx khi lưu kết quả).
- AI đọc báo giá: ai_doc.doc (gói Claude của người dùng), hàng đợi riêng; engine khởi động lại ⇒ đọc tiếp việc dở."""
import os, re, json, uuid, queue, threading, traceback, datetime as dt, unicodedata
import hd_ns as HN

KHOA = threading.Lock(); HANG = queue.Queue(); _CFG = {}
def _kd(s): return unicodedata.normalize("NFD", str(s or "").replace("Đ", "D").replace("đ", "d")).encode("ascii", "ignore").decode()
def _sach(s, n=40): return (re.sub(r"[^A-Za-z0-9]+", "_", _kd(s)).strip("_") or "GOI")[:n]
def ten_gon(s):
    """Tên đơn vị gọn: bỏ phần trong ngoặc / sau dấu phẩy (AI hay chép kèm người đại diện, CCCD, ngày sinh, địa chỉ — thông tin cá nhân không đưa lên bảng)."""
    return re.split(r"\s*[(\[,;]|\s+-\s+|\s+–\s+", str(s or "").strip())[0].strip()[:70]
def so_path(data, da): return os.path.join(data, da, "_HE_THONG", "so_gia.json")
def doc_so(data, da):
    p = so_path(data, da); return json.load(open(p, encoding="utf-8")) if os.path.exists(p) else {}
def _ghi(data, da, s):
    p = so_path(data, da); os.makedirs(os.path.dirname(p), exist_ok=True); t = p + ".tmp"
    json.dump(s, open(t, "w", encoding="utf-8"), ensure_ascii=False, indent=1); os.replace(t, p)
def _sua_dv(data, da, ph, dv, **kv):
    with KHOA:
        s = doc_so(data, da); s[ph]["don_vi"][dv].update(kv); _ghi(data, da, s); return s[ph]["don_vi"][dv]

def ds_phien(data, da):
    out = []
    for p in doc_so(data, da).values():
        dv = p["don_vi"].values()
        out.append(dict(id=p["id"], ten=p["ten"], ngay=p["ngay"], thu_muc=p["thu_muc"], so_dv=len(dv), dang_doc=sum(1 for d in dv if d["trang_thai"] in ("CHO_AI", "DANG_DOC")),
                        chon=(p.get("ket_qua") or {}).get("chon_ten"), luu_luc=(p.get("ket_qua") or {}).get("luc")))
    return sorted(out, key=lambda x: x["ngay"], reverse=True)

def tao_phien(data, da, ten, ghi_chu=""):
    if not str(ten or "").strip(): raise ValueError("Cần tên gói thầu / hạng mục so giá")
    hom_nay = dt.date.today(); i = uuid.uuid4().hex[:10]
    rel = os.path.join(da, "01_THIET_LAP", "03_GOI_THAU", "CHON_THAU", f"{hom_nay:%d.%m.%Y}_{_sach(ten)}")
    if os.path.exists(os.path.join(data, rel)): rel += "_" + i[:4]
    os.makedirs(os.path.join(data, rel, "BAO_GIA"), exist_ok=True)
    with KHOA:
        s = doc_so(data, da)
        s[i] = dict(id=i, ten=ten.strip(), ghi_chu=ghi_chu, ngay=dt.datetime.now().isoformat(timespec="seconds"), thu_muc=rel, don_vi={}, ket_qua=None)
        _ghi(data, da, s)
    return s[i]

def nap(data, da, ph, ten_file, raw):
    s = doc_so(data, da)
    if ph not in s: raise ValueError("Không có phiên so giá này")
    if not ten_file.lower().endswith((".pdf", ".xlsx", ".xlsm")): raise ValueError("Báo giá nhận PDF hoặc Excel (.xlsx)")
    d = os.path.join(data, s[ph]["thu_muc"], "BAO_GIA"); os.makedirs(d, exist_ok=True)
    goc, duoi = os.path.splitext(os.path.basename(ten_file)); f = os.path.join(d, ten_file); n = 2
    while os.path.exists(f): f = os.path.join(d, f"{goc}_{n}{duoi}"); n += 1
    open(f, "wb").write(raw)
    i = uuid.uuid4().hex[:10]
    rec = dict(id=i, ten_file=os.path.basename(f), duong_dan=os.path.relpath(f, data), ten=goc, trang_thai="CHO_AI", ai=None, loi=None, dung=True,
               luc=dt.datetime.now().isoformat(timespec="seconds"))
    with KHOA:
        s = doc_so(data, da); s[ph]["don_vi"][i] = rec; _ghi(data, da, s)
    HANG.put((data, da, ph, i)); return rec

def sua_dv(data, da, ph, dv, ten=None, dung=None):
    kv = {}
    if ten is not None and str(ten).strip(): kv["ten"] = str(ten).strip()[:120]
    if dung is not None: kv["dung"] = bool(dung)
    return _sua_dv(data, da, ph, dv, **kv)

def doc_lai(data, da, ph, dv):
    r = _sua_dv(data, da, ph, dv, trang_thai="CHO_AI", loi=None); HANG.put((data, da, ph, dv)); return r

def _doc_ai(data, da, ph, dv):
    import ai_doc
    rec = doc_so(data, da)[ph]["don_vi"][dv]; _sua_dv(data, da, ph, dv, trang_thai="DANG_DOC")
    try: kq, meta = ai_doc.doc(os.path.join(data, rec["duong_dan"]), cty=_CFG.get("cty", "(chưa khai báo)"))
    except Exception as e: _sua_dv(data, da, ph, dv, trang_thai="LOI_AI", loi=str(e)[:400]); return
    bang = [x for x in HN.NEN.dong_la(kq.get("bang")) if isinstance(x.get("don_gia"), (int, float))]
    canh = []
    if kq.get("loai") not in ("BAO_GIA", "HD_DOI_TAC", "PLHD_DOI_TAC", "GOI_THAU"): canh.append(f"AI nhận file là '{kq.get('loai')}', không phải báo giá — anh kiểm lại file")
    if not bang: canh.append("Không đọc được dòng nào có đơn giá")
    t = kq.get("tong_ghi_tren_file"); tong = sum(HN.NEN.tien_dong(x) for x in HN.NEN.dong_la(kq.get("bang")))
    vat = [v for v in (kq.get("vat_pct"), 0.08, 0.1, 0.05) if isinstance(v, (int, float))]
    if t and tong and abs(tong - t) > t * 0.005 and not any(abs(tong * (1 + v) - t) <= t * 0.002 for v in vat):   # tổng in trên file thường là SAU VAT
        canh.append(f"Σ các dòng {tong:,.0f} ≠ tổng ghi trên file {t:,.0f} — AI có thể đọc sót/thừa dòng")
    ten = rec["ten"] if rec["ten"] != os.path.splitext(rec["ten_file"])[0] else (ten_gon(kq.get("doi_tac_ten")) or rec["ten"])   # anh đã sửa tên thì giữ
    _sua_dv(data, da, ph, dv, trang_thai="XONG", ai=kq, ai_meta=meta, ten=ten, canh_bao=canh, so_dong=len(bang), tong=tong, loi=None,
            luc_ai=dt.datetime.now().isoformat(timespec="seconds"))

def _tho():
    while True:
        data, da, ph, dv = HANG.get()
        try: _doc_ai(data, da, ph, dv)
        except Exception as e: traceback.print_exc(); _sua_dv(data, da, ph, dv, trang_thai="LOI_AI", loi=str(e)[:400])
        finally: HANG.task_done()

def khoi_dong(data, lay_du_an, cty):
    _CFG.update(cty=cty); threading.Thread(target=_tho, daemon=True).start()
    for da in lay_du_an():
        for p in doc_so(data, da).values():
            for d in p["don_vi"].values():
                if d["trang_thai"] in ("CHO_AI", "DANG_DOC"): HANG.put((data, da, p["id"], d["id"]))

def so_sanh(data, da, ph, khung):
    p = doc_so(data, da)[ph]
    dv = [d for d in p["don_vi"].values() if d["dung"] and d["trang_thai"] == "XONG" and d.get("ai")]
    if not dv: raise ValueError("Chưa có báo giá nào AI đọc xong trong phiên này")
    r = HN.ptln(khung, dv); r["phien"] = {k: p[k] for k in ("id", "ten", "ngay", "thu_muc", "ghi_chu")}; r["ket_qua"] = p.get("ket_qua")
    return r

def luu_ket_qua(data, da, ph, khung, chon, ghi_chu=""):
    """CHỈ LƯU kết quả (sổ so_gia.json + PTLN_<gói>_<ngày>.xlsx trong folder phiên). KHÔNG ghi khung."""
    r = so_sanh(data, da, ph, khung); dv = {d["id"]: d for d in r["don_vi"]}
    if chon and chon not in dv: raise ValueError("Đơn vị chọn không nằm trong bảng so sánh")
    luc = dt.datetime.now(); p = doc_so(data, da)[ph]
    f = os.path.join(data, p["thu_muc"], f"PTLN_{_sach(p['ten'])}_{luc:%d.%m.%Y_%H%M}.xlsx"); _xuat_excel(r, f, dv.get(chon), ghi_chu, luc)
    kq = dict(chon=chon, chon_ten=dv[chon]["ten"] if chon else None, ghi_chu=ghi_chu, luc=luc.isoformat(timespec="seconds"), file=os.path.relpath(f, data),
              tom_tat=[{k: d.get(k) for k in ("ten", "hang", "gt_chung", "gt_bao_gia", "vuot_ns", "vuot_ns_tien", "chenh_ns", "so_hang")} for d in r["don_vi"]])
    with KHOA:
        s = doc_so(data, da); s[ph]["ket_qua"] = kq; s[ph].setdefault("lich_su", []).append(kq); _ghi(data, da, s)
    return kq

def _xuat_excel(r, f, chon, ghi_chu, luc):
    """File MỚI (không phải mẫu có sẵn) ⇒ openpyxl an toàn."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "PTLN"; dv = r["don_vi"]
    B, DO, XA = Font(bold=True), PatternFill("solid", fgColor="FDE2E1"), Font(bold=True, color="0F5132")
    ws.append([f"PHÂN TÍCH LỰA CHỌN NHÀ THẦU — {r['phien']['ten']}"]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Lập {luc:%d/%m/%Y %H:%M} · đơn vị chọn: {chon['ten'] if chon else '(chưa chọn)'} · {ghi_chu or ''}"])
    ws.append([]); h = ["Mã NS", "Công việc (ngân sách)", "ĐVT", "KL NS", "ĐG BoQ CĐT", "ĐG NS", "Đã ký thấp nhất"] + [d["ten"] for d in dv] + ["Thấp nhất"]
    ws.append(h); [setattr(c, "font", B) for c in ws[4]]
    for x in r["hang"]:
        ws.append([x["ma_ns"], x["noi_dung"], x["dvt"], x["kl_ns"], x["dg_boq"], x["dg_ns"], x["da_ky_min"]] + [x["dg_dv"].get(str(d["i"]), x["dg_dv"].get(d["i"])) for d in dv] + [x["thap_nhat"]])
        for j, d in enumerate(dv):
            c = ws.cell(ws.max_row, 8 + j); v = c.value
            if isinstance(v, (int, float)) and x["dg_ns"] and v > x["dg_ns"] * 1.005: c.fill = DO
            if isinstance(v, (int, float)) and v == x["thap_nhat"] and len(x["dg_dv"]) > 1: c.font = XA
    ws.append([]); ws.append(["", f"Σ ĐG × KL NS — {r['so_chung']} dòng mọi đơn vị cùng báo", "", "", "", "", ""] + [d["gt_chung"] for d in dv]); ws.cell(ws.max_row, 2).font = B
    ws.append(["", "Hạng", "", "", "", "", ""] + [d.get("hang") for d in dv])
    ws.append(["", "Giá trị báo giá (KL × ĐG của đơn vị)", "", "", "", "", ""] + [d["gt_bao_gia"] for d in dv])
    ws.append(["", "Tiền vượt đơn giá NS", "", "", "", "", ""] + [d["vuot_ns_tien"] for d in dv])
    for row in ws.iter_rows(min_row=5):
        for c in row[3:]:
            if isinstance(c.value, (int, float)): c.number_format = "#,##0.##" if c.column == 4 else "#,##0"
    for j, w in enumerate([14, 46, 7, 12, 13, 12, 14] + [18] * len(dv) + [13], 1): ws.column_dimensions[get_column_letter(j)].width = w
    for c in ws[4]: c.alignment = Alignment(wrap_text=True, vertical="center")
    w2 = wb.create_sheet("Theo_ma_NS"); w2.append(["Mã NS", "Tên", "Ngân sách", "Đã giao (HĐ đã ký)", "NS còn lại"] + [d["ten"] for d in dv]); [setattr(c, "font", B) for c in w2[1]]
    for x in r["theo_ma"]: w2.append([x["ma_ns"], x["ten_ns"], x["ns"], x["da_giao"], x["con_lai"]] + [x["gt"].get(str(d["i"]), x["gt"].get(d["i"])) for d in dv])
    w3 = wb.create_sheet("Chua_ghep_NS"); w3.append(["Đơn vị", "STT", "Nội dung", "ĐVT", "ĐG", "Mã NS", "Lý do"]); [setattr(c, "font", B) for c in w3[1]]
    for l in r["le"]: w3.append([dv[l["i"]]["ten"], l["stt"], l["noi_dung"], l["dvt"], l["dg"], l["ma_ns"], l["ly_do"]])
    w4 = wb.create_sheet("Ghi_chu"); [w4.append([g]) for g in r["ghi_chu"] + ["Nguồn: báo giá AI đọc · N2 ngân sách · N5 BoQ CĐT · N6/N7 HĐ đã ký (giá tham chiếu). File chỉ lưu kết quả, KHÔNG ghi khung."]]
    wb.save(f)
