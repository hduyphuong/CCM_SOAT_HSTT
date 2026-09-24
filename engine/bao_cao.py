"""ĐẦU RA (chỉ đọc) — đọc THẲNG các sheet báo cáo Excel đã tính (R1…R7, 90_Check) + sheet nhập (N1, N4-N9) để có chi tiết.
Không tự tính lại số tài chính: app và Excel luôn ra cùng một con số. Số = lần Excel tính gần nhất (sau mỗi lần ghi sổ COM)."""
import openpyxl, warnings, re
from collections import defaultdict
warnings.filterwarnings("ignore")
LOAI = {"ĐTC": "Đội thi công", "NTP": "Thầu phụ", "NCC": "Nhà cung cấp", "DVK": "Dịch vụ khác", "CĐT": "Chủ đầu tư"}
TRANG = {"DANG_TH": "Đang thực hiện", "DA_QT": "Đã quyết toán", "THANH_LY": "Đã thanh lý", "TAM_DUNG": "Tạm dừng"}
DANG = {"DON_GIA": "Đơn giá cố định", "TRON_GOI": "Trọn gói", "NGUYEN_TAC": "HĐ nguyên tắc (đơn giá)"}
C_DT = dict(ma="A", dt="C", so="E", ngay="F", nd="G", dang="H", gt="I", vat="J", tu="K", tt="L", qt="M", thieu="Y")   # N6: có cột Mã gói
C_CDT = dict(ma="A", dt="C", so="D", ngay="E", nd="F", dang="G", gt="H", vat="I", tu="J", tt="K", qt="L", thieu="X")  # N4

def _n(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0

def _rows(ws, cols, r0=2, key=0):
    """Đọc sheet thành list dict theo danh sách tên cột A, B, C… (bỏ dòng rỗng ở cột khoá)."""
    out = []
    for row in ws.iter_rows(min_row=r0, max_col=len(cols), values_only=True):
        if row[key] in (None, ""): continue
        out.append(dict(zip(cols, row)))
    return out

def doc(path, so_nap=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    w1 = wb["N1_DanhMuc"]
    du_an = dict(ma=w1["A3"].value, ten=w1["B3"].value, cdt=w1["C3"].value, trang_thai=w1["D3"].value, moc=w1["E3"].value)
    doi_tac = {r[6]: dict(ten=r[7] or "", loai=r[8] or "") for r in w1.iter_rows(min_row=3, max_col=20, values_only=True) if r[6]}
    ma_ns = {r[11]: dict(ten=r[12] or "", nhom=r[13] or "") for r in w1.iter_rows(min_row=3, max_col=20, values_only=True) if r[11]}

    # ── R1 CVR: đánh số mục "3.", "3a.", "10." … ──
    cvr = {}
    for a, b, c, d in wb["R1_CVR"].iter_rows(min_row=4, max_col=4, values_only=True):
        m = re.match(r"^(\d+[a-z]?)\.\s*(.*)", str(a or ""))
        if m: cvr[m.group(1)] = dict(ten=m.group(2), nguon=b, gia_tri=c, ghi=d)
    ck = [dict(so=r[0], ten=r[1], trai=r[2], phai=r[3], lech=r[4], ket_luan=str(r[5] or ""), nguon=r[6])
          for r in wb["90_Check"].iter_rows(min_row=2, max_col=7, values_only=True) if r[1]]
    r0_ghi = {r[1]: r[3] for r in wb["R0_TongQuan"].iter_rows(min_row=4, max_col=4, values_only=True) if r[1]}

    # ── R2 ngân sách theo công tác + R7 lãi lỗ (doanh thu phân bổ) ──
    ns = [r for r in _rows(wb["R2_NS_TheoCongTac"], "ma_ns ten nhom ns_r00 dieu_chinh ns cam_ket thuc_hien pct cach con_lai eac ns_tru_eac ns_qt canh_bao ghi".split())
          if r["ma_ns"] in ma_ns]
    w7 = wb["R7_LaiLo_TheoGoi"]; r7, goi, hdr, dong_tong = {}, [], None, {}
    for row in w7.iter_rows(min_row=2, max_col=12, values_only=True):
        a = row[0]
        if a == "Mã gói": hdr = [str(x or "") for x in row]; continue
        if hdr is not None and a:
            goi.append({hdr[i] or f"c{i}": row[i] for i in range(len(hdr))}); continue
        if a in ma_ns: r7[a] = dict(dt_pb=row[3], ln_kh=row[5], ln_dk=row[8], chenh_ln=row[10], canh_bao=row[11], goi=row[2])
        elif a in ("CỘNG MÃ NS", "DT không gắn mã NS", "TỔNG DỰ ÁN"): dong_tong[a] = dict(dt_pb=row[3], ns=row[4], ln_kh=row[5], eac=row[7], ln_dk=row[8])
    # ── HĐ: R3 (số đã tính) + N4/N6 (thông tin HĐ) + R4 (đến hạn) ──
    r3 = {r["ma"]: r for r in _rows(wb["R3_HD_TheoDonVi"], "ma ben ma_dt ten_dt ma_goi gia_tri ns_pb thuc_hien pct tu hu tu_con giu_lai khau_tru tt_lk con_lai trang_thai canh_bao ho_so_thieu".split())}
    r4 = {r["ma"]: r for r in _rows(wb["R4_DenHan"], "ma ben dot_cuoi ngay_cuoi han_tt con_ngay da_tra ngay_nt han_gl giu_lai canh_bao".split())}
    hd = {}
    for sh, C in (("N4_HD_CDT", C_CDT), ("N6_HD_DoiTac", C_DT)):
        w = wb[sh]
        for r in range(2, w.max_row + 1):
            m = w[f"{C['ma']}{r}"].value
            if not m: continue
            g = lambda k: w[f"{C[k]}{r}"].value
            if m in hd: hd[m]["phu_luc"] += 1; continue
            x = r3.get(m, {}); dtc = g("dt"); dang = g("dang")
            hd[m] = dict(ma_hd=m, ben="CĐT" if sh == "N4_HD_CDT" else "ĐỐI TÁC", ma_doi_tac=dtc, doi_tac=doi_tac.get(dtc, {}).get("ten", dtc),
                         loai=LOAI.get(doi_tac.get(dtc, {}).get("loai", ""), doi_tac.get(dtc, {}).get("loai", "")), so_hd=g("so"), ngay_ky=g("ngay"), noi_dung=g("nd"),
                         dang=DANG.get(dang, dang), don_gia=_n(x.get("gia_tri")) == 0, vat=_n(g("vat")), pct_tu=_n(g("tu")), pct_tt=_n(g("tt")), pct_qt=_n(g("qt")),
                         ma_goi=x.get("ma_goi"), gia_tri=_n(x.get("gia_tri")), ns_pb=_n(x.get("ns_pb")), thuc_hien=_n(x.get("thuc_hien")),
                         pct=(_n(x.get("thuc_hien")) / _n(x.get("gia_tri"))) if _n(x.get("gia_tri")) else None,
                         tu_con=_n(x.get("tu_con")), giu_lai=_n(x.get("giu_lai")), khau_tru=_n(x.get("khau_tru")), tt_lk=_n(x.get("tt_lk")), con_lai=_n(x.get("con_lai")),
                         trang_thai=TRANG.get(x.get("trang_thai"), x.get("trang_thai") or ""), canh_bao=x.get("canh_bao") or "",
                         ho_so_thieu=x.get("ho_so_thieu") or g("thieu") or "", den_han=r4.get(m, {}), phu_luc=0, dong=[], dot_cuoi=0)
    lk_kl, lien = defaultdict(float), defaultdict(set)
    for sh in ("N8_TT_CDT", "N9_TT_DoiTac"):
        for r in _rows(wb[sh], "ma dot ngay loai stt nd dvt kl dg tien gt_ky nhom ma_ns tt han da_tra kt nguon".split()):
            if r["loai"] in ("THUC_HIEN", "DIEU_CHINH") and isinstance(r["kl"], (int, float)): lk_kl[(r["ma"], str(r["stt"]))] += r["kl"]
    for sh in ("N5_BOQ_CDT", "N7_HD_DoiTac_ChiTiet"):
        for r in _rows(wb[sh], "ma stt khoa pham_vi nd dvt kl_hd dg thanh_tien ma_cv nhom ma_ns".split()):
            if r["ma"] not in hd: continue
            hd[r["ma"]]["dong"].append(dict(stt=str(r["stt"]), noi_dung=r["nd"], dvt=r["dvt"], kl_hd=r["kl_hd"], don_gia=r["dg"], ngoai=r["pham_vi"] != "TRONG_HD",
                                            kl_lk=lk_kl.get((r["ma"], str(r["stt"])), 0.0)))
            if r["ma_ns"] and hd[r["ma"]]["ben"] != "CĐT": lien[r["ma_ns"]].add(r["ma"])
    # ── BILL: gom N8/N9 theo (HĐ, đợt) — dùng cột ĐÃ TÍNH: giá trị kỳ (K) & tiền thanh toán (N) ──
    bill = {}
    for sh in ("N8_TT_CDT", "N9_TT_DoiTac"):
        for r in _rows(wb[sh], "ma dot ngay loai stt nd dvt kl dg tien gt_ky nhom ma_ns tt han da_tra kt nguon".split()):
            if r["ma"] not in hd or not isinstance(r["dot"], (int, float)) or not r["dot"]: continue
            if r["ma"] not in hd: continue
            k = (r["ma"], int(r["dot"]))
            b = bill.setdefault(k, dict(ma_hd=r["ma"], dot=int(r["dot"]), ngay=None, han=None, da_tra=None, san_luong_ky=0.0, tien_tt=0.0, tam_ung=0.0, hoan_ung=0.0,
                                        so_dong=0, nguon=set(), chua_tra=0))
            if r["loai"] in ("THUC_HIEN", "DIEU_CHINH"): b["san_luong_ky"] += _n(r["gt_ky"])
            if r["loai"] == "TAM_UNG": b["tam_ung"] += _n(r["tien"])
            if r["loai"] == "HOAN_UNG": b["hoan_ung"] += _n(r["tien"])
            b["tien_tt"] += _n(r["tt"]); b["so_dong"] += 1
            for f in ("ngay", "han", "da_tra"):
                if r[f] and (b[f] is None or r[f] > b[f]): b[f] = r[f]
            if not r["da_tra"]: b["chua_tra"] += 1
            if r["nguon"]: b["nguon"].add(str(r["nguon"]).split(" · ")[0])
            hd[r["ma"]]["dot_cuoi"] = max(hd[r["ma"]]["dot_cuoi"], int(r["dot"]))
    qua_app = {(v["phan_loai"].get("ma_hd"), v["tom_tat"].get("dot")) for v in (so_nap or {}).values() if v.get("trang_thai") == "DA_GHI_SO"}
    ds_bill, lk = [], defaultdict(float)
    for (m, dot), b in sorted(bill.items()):
        lk[m] += b["san_luong_ky"]; h = hd[m]
        ds_bill.append(dict(b, nguon=sorted(b["nguon"])[:2], ben=h["ben"], doi_tac=h["doi_tac"], so_hd=h["so_hd"], gia_tri_hd=h["gia_tri"], don_gia=h["don_gia"],
                            san_luong_lk=lk[m], pct=(lk[m] / h["gia_tri"]) if h["gia_tri"] else None, qua_app=(m, dot) in qua_app,
                            tinh_trang="Đã trả" if b["chua_tra"] == 0 else "Chưa ghi ngày trả"))
    for x in ns:
        x.update(r7.get(x["ma_ns"], {}))
        x["hop_dong"] = [dict(ma_hd=m, doi_tac=hd[m]["doi_tac"], so_hd=hd[m]["so_hd"]) for m in sorted(lien.get(x["ma_ns"], []))]
    # ── Đối tác ──
    ds_dt = []
    for ma, d in doi_tac.items():
        hs = [h for h in hd.values() if h["ma_doi_tac"] == ma]
        ds_dt.append(dict(ma=ma, ten=d["ten"], loai=d["loai"], loai_ten=LOAI.get(d["loai"], d["loai"]), so_hd=len(hs), gia_tri=sum(h["gia_tri"] for h in hs),
                          thuc_hien=sum(h["thuc_hien"] for h in hs), tt_lk=sum(h["tt_lk"] for h in hs), giu_lai=sum(h["giu_lai"] for h in hs),
                          ho_so_thieu=sum(1 for h in hs if h["ho_so_thieu"]), ma_hd=[h["ma_hd"] for h in hs]))
    dong_tien = [dict(thang=r["thang"], thu=_n(r["thu"]), chi=_n(r["chi"]), rong=_n(r["rong"]), lk=_n(r["lk"]), ghi=r["ghi"])
                 for r in _rows(wb["R5_DongTien"], "thang tu den thu chi rong lk ghi".split())]
    wd = [r for r in _rows(wb["R6_Workdone"], "nhom ten dvt ma_ns kl_ns kl_cdt kl_cdt_nt kl_doi kl_doi_tt chenh ty_le kl_ns_con dg chi_phi_con co_so canh_bao".split()) if r["canh_bao"]]
    wb.close()
    return dict(du_an=du_an, cvr=cvr, kiem=ck, r0_ghi=r0_ghi, ns=ns, goi=goi, tong_lai_lo=dong_tong, hop_dong=list(hd.values()), bill=ds_bill,
                doi_tac=ds_dt, dong_tien=dong_tien, workdone_canh_bao=wd)
